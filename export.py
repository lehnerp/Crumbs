"""
Export episodes as formal evidence packages.

PDF:  A professional report with one section per episode — timeline,
      source diversity (corroboration), signals/credentials in a table,
      tagged files, and a case metadata header.  Uses reportlab Platypus
      for clean, paged output.

XLSX: Each episode on its own worksheet, signals/credentials as rows.
      A summary sheet lists all episodes with corroboration scores.
      Uses openpyxl with some basic formatting so the cells are readable.
"""
import datetime
import io
import json
import os

# --------------------------------------------------------------------------- #
#  PDF
# --------------------------------------------------------------------------- #

_SEV_COLORS = {"high": (0.85, 0.15, 0.15), "medium": (0.85, 0.5, 0.1), "low": (0.4, 0.65, 0.4)}


def _rl():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, HRFlowable, PageBreak, KeepTogether)
    from reportlab.lib.enums import TA_LEFT
    return (A4, getSampleStyleSheet, ParagraphStyle, mm, colors,
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            HRFlowable, PageBreak, KeepTogether, TA_LEFT)


def build_pdf(episodes, case_meta):
    """Return PDF bytes for the given episodes list."""
    (A4, getSampleStyleSheet, ParagraphStyle, mm, colors,
     SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
     HRFlowable, PageBreak, KeepTogether, TA_LEFT) = _rl()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm,
                            title="Crumbs Episode Evidence",
                            author="Crumbs forensic tool")
    W = A4[0] - 36*mm
    styles = getSampleStyleSheet()
    st_title = ParagraphStyle("ETitle", parent=styles["Title"], fontSize=18, spaceAfter=4)
    st_h1 = ParagraphStyle("EH1", parent=styles["Heading1"], fontSize=13, spaceBefore=10, spaceAfter=3)
    st_h2 = ParagraphStyle("EH2", parent=styles["Heading2"], fontSize=11, spaceBefore=6, spaceAfter=2)
    st_body = ParagraphStyle("EBody", parent=styles["Normal"], fontSize=9, leading=13)
    st_mono = ParagraphStyle("EMono", parent=styles["Normal"], fontSize=8, leading=11,
                             fontName="Courier", wordWrap="LTR")
    st_small = ParagraphStyle("ESmall", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#555555"))
    st_corr = ParagraphStyle("ECorr", parent=styles["Normal"], fontSize=9, leading=12,
                              backColor=colors.HexColor("#e8f4f8"), borderPadding=3)

    now = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    root = case_meta.get("root", "")
    hostname = case_meta.get("hostname", "unknown host")

    story = []

    # ---- cover ----
    story.append(Spacer(1, 20*mm))
    story.append(Paragraph("Episode Evidence Package", st_title))
    story.append(Paragraph(f"<b>Host:</b> {_esc(hostname)}", st_body))
    story.append(Paragraph(f"<b>Evidence folder:</b> {_esc(root)}", st_body))
    story.append(Paragraph(f"<b>Generated:</b> {now}", st_body))
    story.append(Paragraph(f"<b>Episodes included:</b> {len(episodes)}", st_body))
    story.append(Spacer(1, 4*mm))
    story.append(Paragraph(
        "This document is produced automatically by the Crumbs forensic analysis tool. "
        "Each episode is an automatically detected cluster of related events. "
        "Evidence rows link to specific log lines in the collected evidence folder. "
        "Verify each finding against the source log before relying on it.", st_small))
    story.append(PageBreak())

    # ---- episodes ----
    for ep in episodes:
        n = ep.get("n", "?")
        title = ep.get("title", "")
        start = ep.get("start") or ""
        end = ep.get("end") or ""
        minutes = ep.get("minutes", "")
        kind = ep.get("kind", "")
        corr = ep.get("corr_score", 0)
        corr_kinds = ep.get("corr_kinds", [])

        # episode header
        sev_color = colors.HexColor("#cc2222") if corr >= 4 else colors.HexColor("#e07800")
        header_tbl = Table(
            [[Paragraph(f"<b>EP#{n}</b>", ParagraphStyle("ehd", parent=st_body, fontSize=14,
                                                          textColor=colors.white, fontName="Helvetica-Bold")),
              Paragraph(f"<b>{_esc(title)}</b>", ParagraphStyle("etit", parent=st_body, fontSize=12,
                                                                  textColor=colors.white, fontName="Helvetica-Bold"))]],
            colWidths=[22*mm, W-22*mm])
        header_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), sev_color),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(KeepTogether([header_tbl, Spacer(1, 2*mm)]))

        # meta row
        meta_txt = (f"<b>Time:</b> {_esc(start)} → {_esc(end)}"
                    + (f" ({minutes} min)" if minutes else "")
                    + f"  <b>Type:</b> {_esc(kind)}")
        story.append(Paragraph(meta_txt, st_body))

        # corroboration banner
        corr_lbl, corr_sev = _corr_label(corr)
        if corr_lbl:
            bg = {"high": "#d4edda", "medium": "#fff3cd", "low": "#d1ecf1"}.get(corr_sev, "#f8f9fa")
            corr_para = ParagraphStyle("ecp", parent=st_body, fontSize=9, backColor=colors.HexColor(bg),
                                       borderPadding=3)
            sources_txt = ", ".join(corr_kinds[:8]) + ("…" if len(corr_kinds) > 8 else "")
            story.append(Paragraph(
                f"<b>Corroboration:</b> {_esc(corr_lbl)} — confirmed independently by: {_esc(sources_txt)}",
                corr_para))

        story.append(Spacer(1, 2*mm))

        # signals / credentials
        signals = ep.get("signals", [])
        related = ep.get("related", [])
        creds = [s for s in signals if s.get("rule") == "cleartext_creds"]
        other_sigs = [s for s in signals if s.get("rule") != "cleartext_creds"] + related

        if creds:
            story.append(Paragraph("Plaintext credentials captured", st_h2))
            for c in creds[:30]:
                story.append(Paragraph(_esc(c.get("snippet", "")), st_mono))

        if other_sigs:
            story.append(Paragraph(f"Evidence signals ({len(other_sigs)})", st_h2))
            tdata = [["Sev", "Time (UTC)", "Source", "Finding"]]
            for s in other_sigs[:60]:
                sev = s.get("severity", "")
                sc = _SEV_COLORS.get(sev, (0, 0, 0))
                tdata.append([
                    Paragraph(sev, ParagraphStyle("se", parent=st_small, textColor=colors.Color(*sc))),
                    Paragraph(_esc((s.get("ts") or "")[:19]), st_small),
                    Paragraph(_esc(f"{s.get('rel','')}" + (f":{s.get('line','')}" if s.get("line") else "")), st_mono),
                    Paragraph(_esc((s.get("snippet") or s.get("description", ""))[:160]), st_small),
                ])
            sig_tbl = Table(tdata, colWidths=[14*mm, 36*mm, 50*mm, W-100*mm], repeatRows=1)
            sig_tbl.setStyle(_sig_style(colors))
            story.append(sig_tbl)

        # tagged files
        files = ep.get("files", [])
        if files:
            story.append(Spacer(1, 2*mm))
            story.append(Paragraph("Tagged in files", st_h2))
            for f in files:
                story.append(Paragraph(
                    f"<b>{_esc(f.get('rel',''))}:</b> "
                    f"{f.get('events', '?')} event(s), from line {f.get('first_line','?')}",
                    st_body))

        # overlapping episodes
        if ep.get("overlaps"):
            story.append(Spacer(1, 2*mm))
            story.append(Paragraph("Concurrent episodes", st_h2))
            for o in ep["overlaps"]:
                story.append(Paragraph(
                    f"EP#{o['n']}: <b>{_esc(o['title'])}</b> — {_esc(o.get('start',''))} → {_esc(o.get('end',''))}",
                    st_body))

        story.append(Spacer(1, 4*mm))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
        story.append(Spacer(1, 4*mm))

    doc.build(story)
    return buf.getvalue()


def _esc(s):
    if s is None:
        return ""
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _corr_label(score):
    for threshold, lbl, sev in [(6, "confirmed by 6+ independent sources", "high"),
                                  (4, "4–5 independent sources", "medium"),
                                  (2, "2–3 independent sources", "low")]:
        if score >= threshold:
            return lbl, sev
    return "", ""


def _sig_style(colors):
    from reportlab.platypus import TableStyle
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2d2d2d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ])


# --------------------------------------------------------------------------- #
#  XLSX
# --------------------------------------------------------------------------- #

def build_xlsx(episodes, case_meta):
    """Return XLSX bytes for the given episodes list."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # colour palette
    HDR_FILL = PatternFill("solid", fgColor="2D2D2D")
    HDR_FONT = Font(color="FFFFFF", bold=True, size=9)
    EP_FILLS = {"flood": PatternFill("solid", fgColor="D6E4F0"),
                "sniffing": PatternFill("solid", fgColor="FAE5D3"),
                "default": PatternFill("solid", fgColor="EAF4E8")}
    CORR_FILLS = {"high": PatternFill("solid", fgColor="D4EDDA"),
                  "medium": PatternFill("solid", fgColor="FFF3CD"),
                  "low": PatternFill("solid", fgColor="D1ECF1"),
                  "": PatternFill("solid", fgColor="F8F9FA")}
    SEV_FILLS = {"high": PatternFill("solid", fgColor="F8D7DA"),
                 "medium": PatternFill("solid", fgColor="FFF3CD"),
                 "low": PatternFill("solid", fgColor="D4EDDA")}
    thin = Border(left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
                  top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
    wrap = Alignment(wrap_text=True, vertical="top")

    def _hdr(ws, col, text, width=18):
        c = ws.cell(row=1, column=col, value=text)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = wrap
        ws.column_dimensions[get_column_letter(col)].width = width
        return c

    def _cell(ws, row, col, val, fill=None, bold=False, mono=False, wrap_=True):
        c = ws.cell(row=row, column=col, value=val)
        c.border = thin
        c.alignment = Alignment(wrap_text=wrap_, vertical="top")
        if fill:
            c.fill = fill
        if bold:
            c.font = Font(bold=True, size=9)
        elif mono:
            c.font = Font(name="Courier New", size=8)
        else:
            c.font = Font(size=9)
        return c

    # ---- Summary sheet ----
    ws_sum = wb.create_sheet("Summary")
    ws_sum.freeze_panes = "A2"
    for col, (text, w) in enumerate([("EP#", 5), ("Title", 36), ("Type", 14), ("Start", 20),
                                      ("End", 20), ("Minutes", 9), ("Signals", 9),
                                      ("Source files", 9), ("Corroboration", 14), ("Sources", 40)], 1):
        _hdr(ws_sum, col, text, w)

    for row_i, ep in enumerate(episodes, 2):
        n = ep.get("n", "?"); kind = ep.get("kind", "")
        corr = ep.get("corr_score", 0); _, corr_sev = _corr_label(corr)
        corr_fill = CORR_FILLS.get(corr_sev, CORR_FILLS[""])
        ep_fill = EP_FILLS.get(kind, EP_FILLS["default"])
        sigs = len(ep.get("signals", [])) + len(ep.get("related", []))
        lbl, _ = _corr_label(corr)
        _cell(ws_sum, row_i, 1, n, fill=ep_fill, bold=True)
        _cell(ws_sum, row_i, 2, ep.get("title", ""), fill=ep_fill)
        _cell(ws_sum, row_i, 3, kind, fill=ep_fill)
        _cell(ws_sum, row_i, 4, ep.get("start"), fill=ep_fill)
        _cell(ws_sum, row_i, 5, ep.get("end"), fill=ep_fill)
        _cell(ws_sum, row_i, 6, ep.get("minutes"), fill=ep_fill)
        _cell(ws_sum, row_i, 7, sigs, fill=ep_fill)
        _cell(ws_sum, row_i, 8, len(ep.get("files", [])), fill=ep_fill)
        _cell(ws_sum, row_i, 9, f"{corr}: {lbl}" if lbl else str(corr), fill=corr_fill)
        _cell(ws_sum, row_i, 10, ", ".join(ep.get("corr_kinds", [])), fill=corr_fill)

    ws_sum.row_dimensions[1].height = 24

    # ---- Per-episode sheets ----
    for ep in episodes:
        n = ep.get("n", "?"); title = ep.get("title", "")[:28]
        kind = ep.get("kind", "")
        safe_title = title.replace("/", "-").replace("\\", "-").replace(":", "-").replace("?", "").replace("*", "").replace("[", "").replace("]", "")
        ws = wb.create_sheet(f"EP{n} {safe_title}"[:31])
        ws.freeze_panes = "A2"

        # metadata block
        meta_pairs = [
            ("Episode", f"#{n}: {ep.get('title','')}"),
            ("Type", kind), ("Start", ep.get("start", "")), ("End", ep.get("end", "")),
            ("Duration", f"{ep.get('minutes','')} min" if ep.get("minutes") else ""),
            ("Corroboration", f"{ep.get('corr_score',0)} independent sources: {', '.join(ep.get('corr_kinds',[])[:8])}"),
        ]
        if ep.get("overlaps"):
            meta_pairs.append(("Concurrent episodes",
                                ", ".join(f"EP#{o['n']}: {o['title']}" for o in ep["overlaps"])))
        for pair in meta_pairs:
            r = ws.max_row + 1
            c1 = ws.cell(row=r, column=1, value=pair[0])
            c1.font = Font(bold=True, size=9); c1.fill = PatternFill("solid", fgColor="EEEEEE")
            c2 = ws.cell(row=r, column=2, value=pair[1]); c2.font = Font(size=9)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.append([])

        # credentials
        creds = [s for s in ep.get("signals", []) if s.get("rule") == "cleartext_creds"]
        if creds:
            r = ws.max_row + 1
            c = ws.cell(row=r, column=1, value="PLAINTEXT CREDENTIALS CAPTURED")
            c.font = Font(bold=True, color="CC0000", size=9)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            hdr_row = ws.max_row + 1
            for col, (text, w) in enumerate([("Time", 22), ("Source file", 40), ("Snippet", 80)], 1):
                _hdr(ws, col, text, w)
            ws.row_dimensions[hdr_row].height = 18
            for s in creds:
                r2 = ws.max_row + 1
                _cell(ws, r2, 1, s.get("ts"), mono=True)
                _cell(ws, r2, 2, f"{s.get('rel','')}:{s.get('line','')}", mono=True)
                _cell(ws, r2, 3, s.get("snippet", ""))
            ws.append([])

        # signals
        other_sigs = ([s for s in ep.get("signals", []) if s.get("rule") != "cleartext_creds"]
                      + ep.get("related", []))
        if other_sigs:
            r = ws.max_row + 1
            c = ws.cell(row=r, column=1, value=f"Evidence signals ({len(other_sigs)})")
            c.font = Font(bold=True, size=9)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            hdr_row = ws.max_row + 1
            for col, (text, w) in enumerate([("Severity", 10), ("Time (UTC)", 22),
                                              ("Rule", 28), ("Source file", 38), ("Snippet / Description", 70)], 1):
                _hdr(ws, col, text, w)
            ws.row_dimensions[hdr_row].height = 18
            for s in other_sigs[:500]:
                r2 = ws.max_row + 1
                sev = s.get("severity", "")
                sf = SEV_FILLS.get(sev)
                _cell(ws, r2, 1, sev, fill=sf, bold=True)
                _cell(ws, r2, 2, s.get("ts"), mono=True)
                _cell(ws, r2, 3, s.get("rule", s.get("description", "")))
                _cell(ws, r2, 4, f"{s.get('rel','')}:{s.get('line','')}", mono=True)
                _cell(ws, r2, 5, (s.get("snippet") or s.get("description", ""))[:300])
            ws.append([])

        # tagged files
        files = ep.get("files", [])
        if files:
            r = ws.max_row + 1
            c = ws.cell(row=r, column=1, value="Tagged in files")
            c.font = Font(bold=True, size=9)
            hdr_row = ws.max_row + 1
            for col, (text, w) in enumerate([("File", 55), ("Events", 10), ("First line", 12)], 1):
                _hdr(ws, col, text, w)
            for f in files:
                r2 = ws.max_row + 1
                _cell(ws, r2, 1, f.get("rel", ""), mono=True)
                _cell(ws, r2, 2, f.get("events"))
                _cell(ws, r2, 3, f.get("first_line"))

        ws.column_dimensions["A"].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
